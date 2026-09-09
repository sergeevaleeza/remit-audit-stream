"""Matching and workbook-update behaviour, against the synthetic schedule.

These tests guard the two guarantees the app makes to the user: a payment that
is already recorded is never touched twice, and a cell that already holds a
value is never overwritten. Every patient here is fictional (see
`tests/fixtures/README.md`).
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from remit.config import (
    ACTION_FILL,
    ACTION_NEW,
    ACTION_SKIP,
    COL_BILLED,
    COL_COMMENT,
    COL_COPAY,
    COL_PAYMENT,
    SHEET_NAME,
)
from remit.excel_updater import (
    build_updated_workbook,
    download_filename,
    get_schedule_sheet,
    load_schedule_workbook,
    read_schedule_rows,
    resolve_columns,
)
from remit.matching import (
    build_plan,
    is_blank,
    name_score,
    parse_loose_date,
    summarize,
)

from .conftest import find_visit


@pytest.fixture()
def schedule(schedule_bytes):
    """(workbook, worksheet, columns, rows) for the sample schedule."""
    workbook = load_schedule_workbook(schedule_bytes)
    worksheet = get_schedule_sheet(workbook)
    columns = resolve_columns(worksheet)
    return workbook, worksheet, columns, read_schedule_rows(worksheet, columns)


@pytest.fixture()
def plan(visits, schedule):
    return build_plan(visits, schedule[3])


def change_for(plan, patient: str, service_date: date):
    for change in plan:
        if change.visit.patient == patient and change.visit.service_date == service_date:
            return change
    raise AssertionError(f"no planned change for {patient} on {service_date}")


# --- Column resolution -----------------------------------------------------

def test_columns_resolve_by_header_text(schedule):
    _, _, columns, _ = schedule
    assert columns["Patient"] == 1
    assert columns["CPT Code"] == 11
    assert set(columns) >= {"Patient", "Ins", "Data", "Billed", "Payment",
                            "Co-pay", "Comment", "DX", "CPT Code"}


def test_data_starts_on_row_three(schedule):
    _, _, _, rows = schedule
    assert rows[0].row_num == 3
    assert rows[0].patient == "Marlowe, Diane"


# --- Idempotency -----------------------------------------------------------

@pytest.mark.parametrize(
    "patient, service_date",
    [
        ("MARLOWE, DIANE", date(2026, 3, 9)),
        ("MARLOWE, DIANE", date(2026, 4, 21)),
        ("THACKERAY, RENATA", date(2026, 3, 20)),
        ("THACKERAY, RENATA", date(2026, 4, 7)),
        ("THACKERAY, RENATA", date(2026, 4, 24)),
    ],
)
def test_already_paid_visits_are_skipped(plan, patient, service_date):
    change = change_for(plan, patient, service_date)
    assert change.action == ACTION_SKIP
    assert change.fills == {}
    assert not change.accepted


def test_rerunning_the_same_remit_changes_nothing(schedule_bytes, visits, schedule):
    """Second pass over an already-updated workbook must be a no-op."""
    first_plan = build_plan(visits, schedule[3])
    updated, _ = build_updated_workbook(schedule_bytes, first_plan)
    updated_bytes = updated.getvalue()

    workbook = load_schedule_workbook(updated_bytes)
    worksheet = get_schedule_sheet(workbook)
    rows = read_schedule_rows(worksheet, resolve_columns(worksheet))

    second_plan = build_plan(visits, rows)
    assert summarize(second_plan)["fill"] == 0
    assert summarize(second_plan)["new"] == 0

    _, stats = build_updated_workbook(updated_bytes, second_plan)
    assert stats["filled_cells"] == 0
    assert stats["appended_rows"] == 0


def test_skipped_rows_are_untouched_in_the_output(schedule_bytes, plan):
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    # Row 3 = Marlowe 03/09, already paid with a formula in Co-pay.
    assert worksheet.cell(row=3, column=5).value == 166.37
    assert worksheet.cell(row=3, column=6).value == "=23.52+18.92"
    assert worksheet.cell(row=3, column=4).value == "2/32/26"


# --- Fill path -------------------------------------------------------------

def test_fill_path_for_castellano(plan):
    change = change_for(plan, "CASTELLANO, MIGUEL", date(2026, 3, 26))
    assert change.action == ACTION_FILL
    assert change.row_num == 9
    assert change.fills[COL_PAYMENT] == 186.31
    assert change.fills[COL_COPAY] == 47.53
    assert change.fills[COL_COMMENT] == "Dr. A"


def test_fill_writes_only_the_blank_cells(schedule_bytes, plan):
    updated, stats = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    assert worksheet.cell(row=9, column=5).value == 186.31      # Payment
    assert worksheet.cell(row=9, column=6).value == 47.53       # Co-pay
    assert worksheet.cell(row=9, column=9).value == "Dr. A"     # Comment
    assert worksheet.cell(row=9, column=10).value == "F41.1, F32.9"  # DX kept
    assert worksheet.cell(row=9, column=11).value == "99213/90836"   # CPT kept
    assert stats["filled_rows"] >= 1


def test_billed_placeholder_is_never_overwritten(schedule_bytes, plan):
    """`2/32/26` is junk, but it is a value, so the fill must leave it alone."""
    change = change_for(plan, "CASTELLANO, MIGUEL", date(2026, 3, 26))
    assert COL_BILLED not in change.fills
    assert change.existing_billed == "2/32/26"

    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=9, column=4).value == "2/32/26"


def test_blank_billed_would_be_filled(visits, schedule):
    """The never-overwrite rule is about existing values, not about Billed."""
    _, _, _, rows = schedule
    row = next(r for r in rows if r.patient == "Castellano, Miguel")
    row.billed = None
    change = build_plan(
        [find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 3, 26))], rows
    )[0]
    assert change.fills[COL_BILLED] == find_visit(
        visits, "CASTELLANO, MIGUEL", date(2026, 3, 26)
    ).billed_str


# --- Zero is not blank -----------------------------------------------------

def test_zero_payment_counts_as_recorded(plan):
    """Whitfield 03/04/2026 has Payment 0.00 in the sheet."""
    change = change_for(plan, "WHITFIELD, HAROLD", date(2026, 3, 4))
    assert change.action == ACTION_SKIP
    assert change.action != ACTION_FILL


def test_is_blank_treats_zero_and_formulas_as_values():
    assert is_blank(None)
    assert is_blank("")
    assert is_blank("   ")
    assert not is_blank(0)
    assert not is_blank(0.0)
    assert not is_blank("=23.52+18.92")


def test_zero_payment_survives_the_write(schedule_bytes, plan):
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=8, column=5).value == 0.00


# --- New-row path ----------------------------------------------------------

def test_new_row_for_patient_absent_from_the_schedule(plan):
    change = change_for(plan, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    assert change.action == ACTION_NEW
    assert change.row_num is None


def test_new_row_is_appended_with_the_expected_values(schedule_bytes, visits, schedule):
    """One appended row, Ins=Medicare, dates as text, CPT E/M first, DX blank."""
    visit = find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    plan = build_plan([visit], schedule[3])
    updated, stats = build_updated_workbook(schedule_bytes, plan)

    assert stats["appended_rows"] == 1
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    row = 13  # first row after the 10 fixture rows (3..12)
    assert worksheet.cell(row=row, column=1).value == "WHITCOMBE, ROSALIND"
    assert worksheet.cell(row=row, column=2).value == "Medicare"
    assert worksheet.cell(row=row, column=3).value == "05/06/2026"
    assert worksheet.cell(row=row, column=5).value == 166.37
    assert worksheet.cell(row=row, column=6).value == 42.44
    assert worksheet.cell(row=row, column=9).value == "Dr. A"
    assert worksheet.cell(row=row, column=10).value is None      # DX blank
    assert worksheet.cell(row=row, column=11).value == "99213/90833"


def test_new_row_dates_are_written_as_text_not_serials(schedule_bytes, visits, schedule):
    visit = find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    updated, _ = build_updated_workbook(schedule_bytes, build_plan([visit], schedule[3]))
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert isinstance(worksheet.cell(row=13, column=3).value, str)
    assert isinstance(worksheet.cell(row=13, column=4).value, str)


def test_new_rows_do_not_disturb_existing_rows(schedule_bytes, plan):
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=3, column=1).value == "Marlowe, Diane"
    assert worksheet.cell(row=12, column=1).value == "Delacroix, Owen"


# --- Name normalisation ----------------------------------------------------

def test_case_and_whitespace_insensitive_match():
    assert name_score("MARLOWE, DIANE", "Marlowe, Diane") == 100.0
    assert name_score("MARLOWE,  DIANE ", "Marlowe, Diane") == 100.0


def test_truncated_surname_matches(plan):
    """Medicare truncates `Featherstonehaugh` to `FEATHERSTONEH`."""
    change = change_for(plan, "FEATHERSTONEH, WILHELMI", date(2026, 3, 2))
    assert change.action == ACTION_FILL
    assert change.row_num == 10


def test_trailing_middle_initial_is_dropped(plan):
    change = change_for(plan, "SORENSEN, MARCUS R", date(2026, 5, 14))
    assert change.action == ACTION_FILL
    assert change.row_num == 11


def test_different_people_do_not_match():
    assert name_score("CASTELLANO, MIGUEL", "Marlowe, Diane") < 80.0


def test_slavic_feminine_surname_harmonisation():
    assert name_score("IVANOVA, MARIA", "Ivanov, Maria") == 100.0


# --- Free-text date handling ------------------------------------------------

@pytest.mark.parametrize(
    "value, expected",
    [
        ("1/5/2026", date(2026, 1, 5)),
        ("01/27/2026", date(2026, 1, 27)),
        ("3/9/2026", date(2026, 3, 9)),
        ("03/09/26", date(2026, 3, 9)),
        (date(2026, 3, 9), date(2026, 3, 9)),
    ],
)
def test_messy_date_values_parse_by_calendar_value(value, expected):
    assert parse_loose_date(value) == expected


@pytest.mark.parametrize("value", ["2/32/26", "No Billing", "", None, "n/a"])
def test_placeholders_never_parse_as_dates(value):
    assert parse_loose_date(value) is None


def test_dates_match_across_inconsistent_formats(plan):
    """Sheet holds 4/7/2026 as text and 03/20/2026 as a datetime; both match."""
    assert change_for(plan, "THACKERAY, RENATA", date(2026, 4, 7)).action == ACTION_SKIP
    assert change_for(plan, "THACKERAY, RENATA", date(2026, 3, 20)).action == ACTION_SKIP


# --- Workbook integrity ------------------------------------------------------

def test_other_sheets_and_formulas_survive(schedule_bytes, plan):
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    workbook = load_schedule_workbook(updated.getvalue())

    assert workbook.sheetnames == ["2026 Medicare", "2026 Medical"]
    other = workbook["2026 Medical"]
    assert other["A1"].value == "Untouched sheet"
    assert other["A3"].value == "Example, Patient"
    assert other["C3"].value == "=1+1"
    assert other["A1"].font.bold and other["A1"].font.italic


def test_header_and_title_rows_are_preserved(schedule_bytes, plan):
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=1, column=1).value == 2026
    assert worksheet.cell(row=2, column=1).value == "Patient"
    assert worksheet.cell(row=2, column=11).value == "CPT Code"


def test_declined_changes_are_not_written(schedule_bytes, plan):
    """Deselecting a row in the preview must keep it out of the workbook."""
    for change in plan:
        change.accepted = False
    updated, stats = build_updated_workbook(schedule_bytes, plan)
    assert stats["filled_cells"] == 0
    assert stats["appended_rows"] == 0

    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=9, column=5).value is None


def test_only_accepted_changes_are_applied(schedule_bytes, plan):
    for change in plan:
        change.accepted = (
            change.action == ACTION_FILL
            and change.visit.patient == "CASTELLANO, MIGUEL"
        )
    updated, stats = build_updated_workbook(schedule_bytes, plan)
    assert stats["appended_rows"] == 0
    assert stats["filled_rows"] == 1

    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert worksheet.cell(row=9, column=5).value == 186.31
    assert worksheet.cell(row=10, column=5).value is None  # Featherstonehaugh untouched


def test_output_is_a_standalone_workbook(schedule_bytes, plan):
    """The download must open on its own, with no dependency on the app."""
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    reopened = openpyxl.load_workbook(io.BytesIO(updated.getvalue()))
    assert SHEET_NAME in reopened.sheetnames


# --- Summary and filename ----------------------------------------------------

def test_summary_counts_add_up(plan):
    counts = summarize(plan)
    assert counts["visits"] == len(plan)
    assert counts["fill"] + counts["new"] + counts["skip"] + counts["review"] == counts["visits"]


def test_download_filename_is_date_stamped():
    name = download_filename("List_of_Patients_Schedule.xlsx", date(2026, 7, 7))
    assert name == "List_of_Patients_Schedule_updated_2026-07-07.xlsx"

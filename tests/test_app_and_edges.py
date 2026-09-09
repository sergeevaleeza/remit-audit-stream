"""Review paths, workbook repair, and a smoke test of the Streamlit script."""

from __future__ import annotations

import io
import zipfile
from datetime import date

import openpyxl
import pytest

from remit.config import ACTION_FILL, ACTION_NEW, ACTION_REVIEW, ACTION_SKIP, SHEET_NAME
from remit.excel_updater import (
    ScheduleError,
    build_updated_workbook,
    get_schedule_sheet,
    load_schedule_workbook,
    read_schedule_rows,
    repair_workbook_bytes,
    resolve_columns,
)
from remit.matching import Change, ScheduleRow, build_plan, plan_change
from remit.pdf_parser import Visit

from .conftest import find_visit


def make_row(row_num: int, patient: str, data, **overrides) -> ScheduleRow:
    values = dict(patient=patient, data=data, ins="Medicare")
    values.update(overrides)
    return ScheduleRow(row_num=row_num, **values)


def make_visit(patient="TEST, PATIENT", service_date=date(2026, 3, 9),
               npi="1000000001", cpts=("99213", "90833")) -> Visit:
    return Visit(
        patient=patient,
        service_date=service_date,
        npi=npi,
        billed_date=date(2026, 7, 7),
        cpt_codes=list(cpts),
        payment=166.37,
        copay=42.44,
    )


# --- Unknown provider NPI --------------------------------------------------

def test_unknown_npi_is_flagged_and_leaves_comment_blank():
    visit = make_visit(npi="9999999999")
    assert visit.doctor == ""
    assert not visit.known_provider

    change = plan_change(visit, [])
    assert change.action == ACTION_REVIEW
    assert not change.accepted
    assert any("9999999999" in reason for reason in change.review_reasons)


def test_unknown_npi_new_row_omits_comment():
    change = plan_change(make_visit(npi="9999999999"), [])
    assert change.effective_action == ACTION_NEW
    from remit.matching import new_row_values
    assert new_row_values(change.visit)["Comment"] == ""


def test_group_billing_npi_is_not_a_known_doctor():
    from remit.config import GROUP_BILLING_NPI, NPI_TO_DOCTOR
    assert GROUP_BILLING_NPI not in NPI_TO_DOCTOR


# --- Ambiguous matches -----------------------------------------------------

def test_two_unpaid_rows_same_patient_and_date_are_ambiguous():
    rows = [
        make_row(3, "Test, Patient", date(2026, 3, 9)),
        make_row(4, "Test, Patient", date(2026, 3, 9)),
    ]
    change = plan_change(make_visit(), rows)
    assert change.action == ACTION_REVIEW
    assert not change.accepted
    assert any("Ambiguous" in reason for reason in change.review_reasons)


def test_ambiguity_resolved_by_cpt_overlap():
    rows = [
        make_row(3, "Test, Patient", date(2026, 3, 9), cpt="99215/90838"),
        make_row(4, "Test, Patient", date(2026, 3, 9), cpt="99213/90833"),
    ]
    change = plan_change(make_visit(), rows)
    assert change.action == ACTION_FILL
    assert change.row_num == 4


def test_ambiguity_resolved_by_provider_comment():
    rows = [
        make_row(3, "Test, Patient", date(2026, 3, 9), comment="Dr. B"),
        make_row(4, "Test, Patient", date(2026, 3, 9), comment="Dr. A"),
    ]
    change = plan_change(make_visit(), rows)
    assert change.action == ACTION_FILL
    assert change.row_num == 4


def test_ambiguity_resolved_when_only_one_row_is_unpaid():
    rows = [
        make_row(3, "Test, Patient", date(2026, 3, 9), payment=166.37),
        make_row(4, "Test, Patient", date(2026, 3, 9)),
    ]
    change = plan_change(make_visit(), rows)
    assert change.action == ACTION_FILL
    assert change.row_num == 4


def test_accepted_ambiguous_change_actually_writes(schedule_bytes, visits, schedule_rows):
    """Ticking an ambiguous item must have a visible effect, not be a no-op."""
    # A second unpaid row with the same patient, date and CPT as row 9, so
    # neither CPT overlap nor the unpaid check can break the tie.
    rows = schedule_rows + [
        make_row(13, "Castellano, Miguel", date(2026, 3, 26), cpt="99213/90836"),
    ]
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 3, 26))
    change = plan_change(visit, rows)
    assert change.action == ACTION_REVIEW
    assert change.fills, "an ambiguous match should still carry fill values"

    change.accepted = True
    _, stats = build_updated_workbook(schedule_bytes, [change])
    assert stats["filled_cells"] > 0


# --- Low-confidence names --------------------------------------------------

def test_low_confidence_name_needs_review():
    """A couple of typos land in the review band (~91%), not auto-matched."""
    rows = [make_row(3, "Testname, Sammpel", date(2026, 3, 9))]
    change = plan_change(make_visit(patient="TESTNAME, SAMPLE"), rows)
    assert change.action == ACTION_REVIEW
    assert not change.accepted
    assert any("Low-confidence" in reason for reason in change.review_reasons)


def test_truncation_is_treated_as_a_confident_match_not_a_typo():
    """A strict prefix is how Medicare truncates, so it matches outright.

    The trade-off is that a dropped trailing character reads as truncation
    rather than as a typo; the exact-date requirement is what keeps this safe.
    """
    from remit.matching import name_score
    assert name_score("TESTNAME, SAMPLE", "Testnamee, Samplee") == 100.0


def test_unrelated_name_becomes_a_new_row_not_a_review():
    rows = [make_row(3, "Zimmerman, Robert", date(2026, 3, 9))]
    change = plan_change(make_visit(patient="NOBODY, HERE"), rows)
    assert change.action == ACTION_NEW


def test_matching_name_on_a_different_date_is_a_new_row():
    rows = [make_row(3, "Test, Patient", date(2026, 1, 1))]
    change = plan_change(make_visit(service_date=date(2026, 3, 9)), rows)
    assert change.action == ACTION_NEW


def test_effective_action_derivation():
    review_fill = Change(visit=make_visit(), action=ACTION_REVIEW, row_num=5,
                         fills={"Payment": 1.0})
    review_new = Change(visit=make_visit(), action=ACTION_REVIEW)
    assert review_fill.effective_action == ACTION_FILL
    assert review_new.effective_action == ACTION_NEW
    assert Change(visit=make_visit(), action=ACTION_SKIP).effective_action == ACTION_SKIP


def test_review_items_are_not_applied_unless_accepted(schedule_bytes):
    change = plan_change(make_visit(npi="9999999999"), [])
    assert not change.accepted
    _, stats = build_updated_workbook(schedule_bytes, [change])
    assert stats["appended_rows"] == 0


# --- Workbook repair -------------------------------------------------------

def _rewrite_font_family(data: bytes, value: bytes) -> bytes:
    """Inject an out-of-range font family, as real Excel files contain."""
    source = io.BytesIO(data)
    target = io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w") as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                payload = payload.replace(b'<family val="2"/>', b'<family val="' + value + b'"/>')
            zout.writestr(item, payload)
    return target.getvalue()


def test_openpyxl_rejects_out_of_range_font_family(schedule_bytes):
    """Guards the assumption behind the repair: this really does fail."""
    broken = _rewrite_font_family(schedule_bytes, b"34")
    with pytest.raises(ValueError):
        openpyxl.load_workbook(io.BytesIO(broken))


def test_repair_allows_the_workbook_to_load(schedule_bytes):
    broken = _rewrite_font_family(schedule_bytes, b"34")
    workbook = load_schedule_workbook(broken)
    assert SHEET_NAME in workbook.sheetnames
    assert workbook[SHEET_NAME].cell(row=3, column=1).value == "Marlowe, Diane"


def test_repair_preserves_every_sheet_and_cell(schedule_bytes):
    repaired = repair_workbook_bytes(_rewrite_font_family(schedule_bytes, b"18"))
    original = openpyxl.load_workbook(io.BytesIO(schedule_bytes))
    fixed = openpyxl.load_workbook(io.BytesIO(repaired))

    assert fixed.sheetnames == original.sheetnames
    for name in original.sheetnames:
        for row in range(1, original[name].max_row + 1):
            for col in range(1, original[name].max_column + 1):
                assert fixed[name].cell(row=row, column=col).value == \
                       original[name].cell(row=row, column=col).value


def test_valid_workbook_is_not_rewritten(schedule_bytes):
    """The repair path only runs on failure, leaving good files untouched."""
    workbook = load_schedule_workbook(schedule_bytes)
    assert workbook[SHEET_NAME].cell(row=2, column=1).value == "Patient"


# --- Workbook validation ---------------------------------------------------

def test_missing_sheet_is_reported_clearly():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Something Else"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ScheduleError, match=SHEET_NAME):
        get_schedule_sheet(load_schedule_workbook(buffer.getvalue()))


def test_missing_header_is_reported_clearly():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet["A1"] = 2026
    for index, header in enumerate(["Patient", "Ins", "Data"], start=1):
        sheet.cell(row=2, column=index, value=header)
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ScheduleError, match="Payment"):
        resolve_columns(get_schedule_sheet(load_schedule_workbook(buffer.getvalue())))


def test_columns_are_found_even_when_reordered():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet["A1"] = 2026
    shuffled = ["CPT Code", "Patient", "Payment", "Data", "Co-pay",
                "Billed", "Comment", "Ins", "DX", "Office", "Co-pays Paid"]
    for index, header in enumerate(shuffled, start=1):
        sheet.cell(row=2, column=index, value=header)
    sheet.cell(row=3, column=2, value="Test, Patient")

    columns = resolve_columns(sheet)
    assert columns["Patient"] == 2
    assert columns["CPT Code"] == 1
    assert columns["Payment"] == 3


# --- Streamlit script smoke test ------------------------------------------

def test_streamlit_script_runs_without_error():
    """The app script must import and execute to its first stop() cleanly."""
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file("streamlit_app.py", default_timeout=60).run()
    assert not app.exception, [str(e) for e in app.exception]
    assert any("Upload the schedule workbook" in str(info.value) for info in app.info)

"""The optional `AMSMC_employees.xlsx` upload, and the POS -> `Ins` rule.

Every patient here is fictional (see `tests/fixtures/README.md`). The synthetic
employees workbook reproduces the real file's awkward parts: Oxana's header sits
on row 2 while Ana's and Marcia's sit on row 1, her patient column is headed
`Patient` rather than `Patient Name`, and several labels carry doubled internal
spaces.

The tabs are populated from the reconciled EOB visits, never from the schedule
-- see `test_employees_from_eob.py` for the rules that decide which tab gets
what.
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from remit.config import (
    EMP_ACTION_APPEND,
    EMP_ACTION_FILL,
    EMP_ACTION_MISMATCH,
    EMP_ACTION_UNASSIGNED,
    EMPLOYEE_PAYMENT_COLUMN,
    INSURANCE_VALUE,
    TELEHEALTH_INSURANCE_LABEL,
)
from remit.employees import (
    EmployeesError,
    build_updated_employees,
    column_for,
    employees_download_filename,
    eob_visits,
    find_header_row,
    header_key,
    load_employees_workbook,
    plan_employee_changes,
    read_sheets,
    summarize_employees,
)
from remit.excel_updater import (
    get_schedule_sheet,
    load_schedule_workbook,
    read_schedule_rows,
    resolve_columns,
)
from remit.matching import build_plan
from remit.pdf_parser import insurance_label, parse_service_line

from .conftest import find_visit

STAMP = date(2026, 9, 10)


@pytest.fixture()
def sheets(employee_sheets):
    return employee_sheets


@pytest.fixture()
def emp_changes(sheets, all_visits):
    return plan_employee_changes(sheets, eob_visits(all_visits), today=STAMP)


def change_for(changes, patient, session_date, sheet=None):
    for change in changes:
        if (change.patient.startswith(patient.split(",")[0])
                and change.session_date == session_date
                and (sheet is None or change.sheet == sheet)):
            return change
    raise AssertionError(f"no employee change for {patient} on {session_date}")


def sheet_values(worksheet, header_row):
    """{(patient, date): {header: value}} for one provider tab."""
    headers = {
        c: worksheet.cell(row=header_row, column=c).value
        for c in range(1, worksheet.max_column + 1)
    }
    patient_col = next(c for c, h in headers.items()
                       if header_key(h) in ("patient name", "patient"))
    date_col = next(c for c, h in headers.items() if header_key(h) == "date of session")

    out = {}
    for row in range(header_row + 1, worksheet.max_row + 1):
        patient = worksheet.cell(row=row, column=patient_col).value
        if patient in (None, ""):
            continue
        key = (str(patient), str(worksheet.cell(row=row, column=date_col).value))
        out[key] = {
            headers[c]: worksheet.cell(row=row, column=c).value
            for c in headers if headers[c]
        }
    return out


# --- POS / telehealth -------------------------------------------------------

def test_pos_and_modifier_are_captured():
    line = parse_service_line(
        "1000000001 0313 031326 10 1.0 99213 95 200.00 117.58 0.00 23.52 CO-45 82.42 92.18",
        "TEST, PATIENT", "f.pdf",
    )
    assert line.pos == "10"
    assert line.modifiers == ("95",)
    assert line.telehealth


def test_office_visit_is_not_telehealth():
    line = parse_service_line(
        "1000000001 0309 030926 11 1.0 99213 200.00 117.58 0.00 23.52 CO-45 82.42 92.18",
        "TEST, PATIENT", "f.pdf",
    )
    assert line.pos == "11"
    assert line.modifiers == ()
    assert not line.telehealth


def test_pos_10_with_95_yields_the_telehealth_label(visits):
    """`FEATHERSTONEH, WILHELMI` 03/02/2026 is billed POS 10 + 95."""
    visit = find_visit(visits, "FEATHERSTONEH, WILHELMI", date(2026, 3, 2))
    assert visit.telehealth
    assert insurance_label(visit) == TELEHEALTH_INSURANCE_LABEL == "POS 10(95)"
    assert visit.insurance == "POS 10(95)"


def test_pos_11_yields_medicare(visits):
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 3, 26))
    assert not visit.telehealth
    assert insurance_label(visit) == INSURANCE_VALUE == "Medicare"


def test_pos_10_without_the_modifier_is_not_telehealth():
    line = parse_service_line(
        "1000000001 0313 031326 10 1.0 99213 200.00 117.58 0.00 23.52 CO-45 82.42 92.18",
        "TEST, PATIENT", "f.pdf",
    )
    assert not line.telehealth


def test_new_schedule_row_takes_ins_from_the_visit(visits):
    """A telehealth visit appended to the schedule says `POS 10(95)`."""
    from remit.matching import new_row_values

    visit = find_visit(visits, "PETROSSIAN, KNARIK S", date(2025, 12, 25))
    assert new_row_values(visit)["Ins"] == "POS 10(95)"


def test_existing_medicare_row_is_flagged_not_overwritten(schedule_bytes, visits):
    """Never-overwrite still holds: the row is flagged, `Ins` left alone."""
    worksheet = get_schedule_sheet(load_schedule_workbook(schedule_bytes))
    rows = read_schedule_rows(worksheet, resolve_columns(worksheet))
    plan = build_plan(visits, rows, today=STAMP)

    flagged = [c for c in plan if c.telehealth_mismatch]
    assert flagged
    for change in flagged:
        assert "left as-is" in change.telehealth_mismatch
        assert change.ins_update == ""


def test_telehealth_label_reaches_the_employees_visit(all_visits):
    petrossian = next(v for v in eob_visits(all_visits)
                      if v.patient.startswith("Petrossian"))
    assert petrossian.insurance == "POS 10(95)"


# --- Reading the workbook ---------------------------------------------------

def test_header_row_found_per_sheet(sheets):
    """Ana and Marcia on row 1, Oxana on row 2."""
    assert sheets["Ana"].header_row == 1
    assert sheets["Marcia"].header_row == 1
    assert sheets["Oxana"].header_row == 2


def test_oxana_patient_column_is_headed_patient(sheets):
    oxana = sheets["Oxana"]
    assert column_for(oxana.columns, "Patient") == oxana.patient_column
    assert oxana.patient_column == 1


def test_headers_match_despite_doubled_spaces(sheets):
    """`Co-payment   Old` must resolve like `Co-payment Old`."""
    assert column_for(sheets["Ana"].columns, "Co-payment Old") is not None
    assert header_key("Co-payment   Old") == "co-payment old"


def test_payment_column_differs_per_provider(sheets):
    assert EMPLOYEE_PAYMENT_COLUMN["Ana"] == "Paid by Ins toAna"
    assert EMPLOYEE_PAYMENT_COLUMN["Oxana"] == "Paid by Insurance"
    assert sheets["Ana"].payment_column != sheets["Ana"].columns[header_key("Paid by Insurance")]
    assert sheets["Oxana"].payment_column == sheets["Oxana"].columns[header_key("Paid by Insurance")]


def test_non_provider_sheets_are_ignored(sheets):
    assert set(sheets) == {"Ana", "Marcia", "Oxana"}


def test_workbook_without_provider_sheets_is_rejected():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Something Else"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(EmployeesError, match="provider sheets"):
        read_sheets(load_employees_workbook(buffer.getvalue()))


def test_sheet_without_a_header_row_is_rejected():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ana"
    sheet["A1"] = "nothing recognisable"
    with pytest.raises(EmployeesError):
        find_header_row(sheet)


def test_has_session_is_the_append_dedup_key(sheets):
    ana = sheets["Ana"]
    assert ana.has_session("Castellano, Miguel", date(2026, 3, 26))
    assert not ana.has_session("Castellano, Miguel", date(2026, 4, 21))
    assert not ana.has_session("Nobody, Atall", date(2026, 3, 26))


# --- Routing by name + date -------------------------------------------------

def test_existing_row_is_filled_by_name_and_date(emp_changes):
    change = change_for(emp_changes, "Castellano", "03/26/2026", sheet="Ana")
    assert change.action == EMP_ACTION_FILL
    assert change.payment == 186.31
    assert change.payment_column == "Paid by Ins toAna"


def test_a_supervising_npi_never_blocks_a_fill(emp_changes):
    """Marcia's sessions bill incident-to, which is not a conflict."""
    change = change_for(emp_changes, "Whitfield", "03/04/2026", sheet="Marcia")
    assert change.action == EMP_ACTION_FILL


def test_oxana_row_fills_normally(emp_changes):
    change = change_for(emp_changes, "Thackeray", "04/07/2026", sheet="Oxana")
    assert change.action == EMP_ACTION_FILL


def test_another_associates_npi_is_a_mismatch(emp_changes):
    """`Beaumont` sits in Ana's tab but Oxana's NPI performed the session."""
    change = change_for(emp_changes, "Beaumont", "06/10/2026", sheet="Ana")
    assert change.action == EMP_ACTION_MISMATCH
    assert not change.accepted
    assert "Oxana" in change.note


def test_suffix_aware_matching_still_applies(sheets, all_visits):
    """`BYSTRITSKAYA, ANNA` must match a `Bystritskaya Jr, Anna` tab row."""
    from remit.employees import _matches_for

    matches = _matches_for(eob_visits(all_visits), "Bystritskaya Jr, Anna",
                           date(2026, 4, 14))
    assert len(matches) == 1


# --- Append behaviour -------------------------------------------------------

def test_new_session_appends_for_ana(emp_changes):
    """Only because 06/16 was billed under Ana's own NPI."""
    change = change_for(emp_changes, "Ravensworth", "06/16/2026", sheet="Ana")
    assert change.action == EMP_ACTION_APPEND


def test_new_session_appends_for_oxana(emp_changes):
    change = change_for(emp_changes, "Nakamura", "06/17/2026", sheet="Oxana")
    assert change.action == EMP_ACTION_APPEND


def test_marcia_never_appends(emp_changes):
    """Her tab roster defines ownership; the remit cannot add to it."""
    assert not [c for c in emp_changes if c.sheet == "Marcia" and c.appends]


def test_patient_in_no_tab_is_unassigned(emp_changes):
    change = change_for(emp_changes, "Okafor", "04/28/2026")
    assert change.action == EMP_ACTION_UNASSIGNED
    assert change.sheet == ""
    assert not change.accepted


def test_unassigned_visits_are_never_placed_by_guesswork(emp_changes):
    unassigned = [c for c in emp_changes if c.action == EMP_ACTION_UNASSIGNED]
    assert unassigned
    assert all(c.sheet == "" and not c.fills and not c.accepted for c in unassigned)


# --- Writing ----------------------------------------------------------------

def test_apply_fills_and_appends(employees_bytes, emp_changes):
    updated, stats = build_updated_employees(employees_bytes, emp_changes)
    expected_appends = sum(1 for c in emp_changes if c.accepted and c.appends)
    assert stats["filled_rows"] > 0
    assert stats["appended_rows"] == expected_appends > 0

    workbook = load_employees_workbook(updated.getvalue())
    ana = sheet_values(workbook["Ana"], 1)
    row = ana[("Castellano, Miguel", "03/26/2026")]
    assert row["Insurance"] == "Medicare"
    assert row["Co-pay by EOB"] == 47.53
    assert row["Paid by Ins toAna"] == 186.31


def test_oxana_mapping_uses_paid_by_insurance(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    workbook = load_employees_workbook(updated.getvalue())
    row = sheet_values(workbook["Oxana"], 2)[("Thackeray, Renata", "04/07/2026")]
    assert row["Paid by Insurance"] == 186.31
    assert row["Insurance"] == "Medicare"


def test_appended_row_carries_the_mapped_columns(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    workbook = load_employees_workbook(updated.getvalue())
    row = sheet_values(workbook["Ana"], 1)[("Ravensworth, Cecily", "06/16/2026")]
    assert row["Insurance"] == "Medicare"
    assert row["Co-pay by EOB"] == 47.53
    assert row["Paid by Ins toAna"] == 186.31
    # Unmapped columns stay blank.
    assert row["Memo"] is None
    assert row["Billed"] is None


def test_never_overwrites_a_populated_cell(employees_bytes, emp_changes):
    """Castellano 04/09 already has 999.99 recorded in Ana's tab."""
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    workbook = load_employees_workbook(updated.getvalue())
    row = sheet_values(workbook["Ana"], 1)[("Castellano, Miguel", "04/09/2026")]
    assert row["Paid by Ins toAna"] == 999.99


def test_mismatched_row_is_not_filled(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    workbook = load_employees_workbook(updated.getvalue())
    row = sheet_values(workbook["Ana"], 1)[("Beaumont, Sylvie", "06/10/2026")]
    assert row["Paid by Ins toAna"] is None
    assert row["Insurance"] is None


def test_appended_rows_inherit_styling(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    worksheet = load_employees_workbook(updated.getvalue())["Ana"]
    # Ana's tab has six seeded rows (2-7); the append lands on row 8.
    template, appended = 7, 8
    assert worksheet.cell(row=appended, column=1).value == "Ravensworth, Cecily"
    for column in (1, 2, 8):
        assert (worksheet.cell(row=appended, column=column).number_format
                == worksheet.cell(row=template, column=column).number_format)
        assert (worksheet.cell(row=appended, column=column).font.name
                == worksheet.cell(row=template, column=column).font.name)


def test_rerun_appends_nothing_new(employees_bytes, sheets, all_visits):
    """Dedup by patient + Date of Session: no double-adds."""
    visits = eob_visits(all_visits)
    first = plan_employee_changes(sheets, visits, today=STAMP)
    once, _ = build_updated_employees(employees_bytes, first)

    reread = read_sheets(load_employees_workbook(once.getvalue()))
    second = plan_employee_changes(reread, visits, today=STAMP)
    _, stats = build_updated_employees(once.getvalue(), second)

    assert stats["appended_rows"] == 0
    assert stats["filled_cells"] == 0


def test_every_sheet_survives(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    workbook = load_employees_workbook(updated.getvalue())
    assert workbook.sheetnames == ["Ana", "Marcia", "Oxana", "Notes"]
    assert workbook["Notes"]["A1"].value == "Not a provider tab"


def test_declined_changes_write_nothing(employees_bytes, emp_changes):
    for change in emp_changes:
        change.accepted = False
    _, stats = build_updated_employees(employees_bytes, emp_changes)
    assert stats["filled_cells"] == 0
    assert stats["appended_rows"] == 0


def test_output_is_a_standalone_workbook(employees_bytes, emp_changes):
    updated, _ = build_updated_employees(employees_bytes, emp_changes)
    assert openpyxl.load_workbook(io.BytesIO(updated.getvalue()))


# --- The EOB-side view ------------------------------------------------------

def test_eob_visits_are_title_cased(all_visits):
    """Medicare shouts; the tabs do not."""
    castellano = next(v for v in eob_visits(all_visits)
                      if v.patient.startswith("Castellano"))
    assert castellano.patient == "Castellano, Miguel"


def test_eob_visits_are_unique_per_patient_date_and_npi(all_visits):
    keys = [(v.patient.lower(), v.session_date, v.npi) for v in eob_visits(all_visits)]
    assert len(keys) == len(set(keys))


def test_eob_visit_amounts_come_straight_from_the_remit(all_visits):
    castellano = next(v for v in eob_visits(all_visits)
                      if v.patient.startswith("Castellano")
                      and v.session_date == date(2026, 3, 26))
    assert castellano.payment == 186.31
    assert castellano.copay == 47.53


# --- Summary / filename -----------------------------------------------------

def test_summary_counts(emp_changes):
    counts = summarize_employees(emp_changes)
    assert counts["rows"] == len(emp_changes)
    assert counts["fill"] > 0
    assert counts["append"] == sum(
        1 for c in emp_changes if c.action == EMP_ACTION_APPEND
    ) > 0
    assert counts["mismatch"] >= 1
    assert counts["ambiguous"] >= 1


def test_download_filename_has_no_patient_name():
    name = employees_download_filename("AMSMC_employees.xlsx", date(2026, 9, 10))
    assert name == "AMSMC_employees_updated_2026-09-10.xlsx"


def test_employees_upload_is_optional(schedule_bytes, visits):
    """Without the workbook the app behaves exactly as before."""
    worksheet = get_schedule_sheet(load_schedule_workbook(schedule_bytes))
    rows = read_schedule_rows(worksheet, resolve_columns(worksheet))
    plan = build_plan(visits, rows, today=STAMP)
    assert plan  # the schedule side is unaffected by the absent workbook

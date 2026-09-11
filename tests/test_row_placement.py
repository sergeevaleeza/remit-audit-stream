"""New rows group under their patient instead of always appending."""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from remit.config import ACTION_NEW, DATA_START_ROW, SHEET_NAME
from remit.excel_updater import (
    build_updated_workbook,
    get_schedule_sheet,
    insertion_blocked_reason,
    load_schedule_workbook,
    read_schedule_rows,
    resolve_columns,
)
from remit.matching import build_plan, find_patient_anchor, plan_change

from .conftest import find_row, find_visit, sheet_rows

PATIENT_COL = 1
DATA_COL = 3
PAYMENT_COL = 5


def patients_in_order(worksheet) -> list[str]:
    return [str(row["Patient"]) for row in sheet_rows(worksheet)]


def applied_sheet(schedule_bytes, plan):
    updated, stats = build_updated_workbook(schedule_bytes, plan)
    return load_schedule_workbook(updated.getvalue())[SHEET_NAME], stats


# --- Anchor selection ------------------------------------------------------

def test_anchor_is_the_patients_last_existing_row(visits, schedule_rows):
    """`Castellano, Miguel` occupies row 9 in the fixture."""
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    anchor, name = find_patient_anchor(visit, schedule_rows)
    assert anchor == 9
    assert name == "Castellano, Miguel"


def test_anchor_is_none_for_an_unknown_patient(visits, schedule_rows):
    visit = find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    assert find_patient_anchor(visit, schedule_rows) == (None, None)


def test_anchor_uses_the_last_of_several_rows(visits, schedule_rows):
    """`Marlowe, Diane` has rows 3 and 4; the anchor must be 4."""
    visit = find_visit(visits, "MARLOWE, DIANE", date(2026, 3, 9))
    anchor, _ = find_patient_anchor(visit, schedule_rows)
    assert anchor == 4


def test_anchor_uses_the_last_occurrence_when_rows_are_non_contiguous(visits, schedule_rows):
    """A patient split across the sheet anchors under their *last* row."""
    from remit.matching import ScheduleRow

    rows = list(schedule_rows) + [ScheduleRow(row_num=99, patient="Castellano, Miguel",
                                              data=date(2026, 1, 2))]
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    anchor, _ = find_patient_anchor(visit, rows)
    assert anchor == 99


def test_suffix_name_anchors_under_the_suffix_row(visits, schedule_rows):
    """`BYSTRITSKAYA, ANNA` must anchor under `Bystritskaya Jr, Anna`."""
    visit = find_visit(visits, "BYSTRITSKAYA, ANNA", date(2026, 4, 14))
    anchor, name = find_patient_anchor(visit, schedule_rows)
    assert name == "Bystritskaya Jr, Anna"
    assert anchor == 12


# --- Placement in the written sheet ----------------------------------------

def test_new_visit_is_inserted_directly_below_its_patient(schedule_bytes, visits, schedule_rows):
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    worksheet, stats = applied_sheet(schedule_bytes, build_plan([visit], schedule_rows))

    assert stats["inserted_rows"] == 1
    assert stats["appended_rows"] == 0
    assert worksheet.cell(row=9, column=PATIENT_COL).value == "Castellano, Miguel"
    assert worksheet.cell(row=10, column=PATIENT_COL).value == "Castellano, Miguel"
    assert worksheet.cell(row=10, column=DATA_COL).value == "04/09/2026"
    # Everything below shifted down by exactly one.
    assert worksheet.cell(row=11, column=PATIENT_COL).value == "Featherstonehaugh, Wilhelmina"


def test_several_new_visits_insert_as_one_date_ordered_block(schedule_bytes, visits, schedule_rows):
    castellano = [
        find_visit(visits, "CASTELLANO, MIGUEL", d)
        for d in (date(2026, 5, 7), date(2026, 4, 9), date(2026, 4, 21))
    ]
    worksheet, stats = applied_sheet(schedule_bytes, build_plan(castellano, schedule_rows))

    assert stats["inserted_rows"] == 3
    assert [worksheet.cell(row=r, column=DATA_COL).value for r in (10, 11, 12)] == [
        "04/09/2026", "04/21/2026", "05/07/2026",
    ]
    assert all(
        worksheet.cell(row=r, column=PATIENT_COL).value == "Castellano, Miguel"
        for r in (9, 10, 11, 12)
    )


def test_unknown_patient_is_appended_at_the_bottom(schedule_bytes, visits, schedule_rows):
    visit = find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    worksheet, stats = applied_sheet(schedule_bytes, build_plan([visit], schedule_rows))

    assert stats["appended_rows"] == 1
    assert stats["inserted_rows"] == 0
    assert worksheet.cell(row=worksheet.max_row, column=PATIENT_COL).value == "Whitcombe, Rosalind"


def test_suffix_patient_groups_instead_of_appending(schedule_bytes, visits, schedule_rows):
    """A new visit for `BYSTRITSKAYA, ANNA` joins `Bystritskaya Jr, Anna`."""
    from remit.pdf_parser import Visit

    visit = Visit(
        patient="BYSTRITSKAYA, ANNA",
        service_date=date(2026, 6, 2),
        npi="1000000001",
        billed_date=date(2026, 8, 14),
        cpt_codes=["99213", "90833"],
        payment=166.37,
        copay=42.44,
        check_efts={"900000001"},
    )
    change = plan_change(visit, schedule_rows)
    assert change.effective_action == ACTION_NEW
    assert change.inserts_under_patient

    worksheet, stats = applied_sheet(schedule_bytes, [change])
    assert stats["inserted_rows"] == 1
    assert stats["appended_rows"] == 0
    assert worksheet.cell(row=12, column=PATIENT_COL).value == "Bystritskaya Jr, Anna"
    assert worksheet.cell(row=13, column=PATIENT_COL).value == "Bystritskaya, Anna"
    assert worksheet.cell(row=13, column=DATA_COL).value == "06/02/2026"


def test_inserts_and_appends_together(schedule_bytes, visits, schedule_rows):
    """Full plan: known patients group, unknown ones land at the bottom."""
    plan = build_plan(visits, schedule_rows)
    worksheet, stats = applied_sheet(schedule_bytes, plan)

    # Grouped under an existing patient vs appended at the bottom.
    grouped = [c for c in plan if c.accepted and c.effective_action == ACTION_NEW
               and c.inserts_under_patient]
    bottom = [c for c in plan if c.accepted and c.effective_action == ACTION_NEW
              and not c.inserts_under_patient]
    assert stats["inserted_rows"] == len(grouped) > 0
    assert stats["appended_rows"] == len(bottom) > 0

    order = patients_in_order(worksheet)
    assert order[:8] == [
        "Marlowe, Diane", "Marlowe, Diane",
        "Thackeray, Renata", "Thackeray, Renata", "Thackeray, Renata",
        # Whitfield's second session groups under his existing row.
        "Whitfield, Harold", "Whitfield, Harold",
        "Castellano, Miguel",
    ]
    assert order[-4:] == [
        "Okafor, Chidinma",
        "Petrossian, Knarik S", "Petrossian, Knarik S",
        "Whitcombe, Rosalind",
    ]


# --- Integrity -------------------------------------------------------------

def test_row_count_grows_by_exactly_the_new_rows(schedule_bytes, visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    worksheet, stats = applied_sheet(schedule_bytes, plan)

    before = len(schedule_rows)
    after = len(sheet_rows(worksheet))
    assert after == before + stats["new_rows"]
    assert stats["new_rows"] == stats["inserted_rows"] + stats["appended_rows"]


def test_no_existing_row_is_lost_or_duplicated(schedule_bytes, visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    worksheet, _ = applied_sheet(schedule_bytes, plan)

    for original in schedule_rows:
        found = find_row(worksheet, str(original.patient), original.data)
        assert found["CPT Code"] == original.cpt
        assert found["Ins"] == original.ins


def test_existing_values_are_unchanged_for_untouched_rows(schedule_bytes, visits, schedule_rows):
    """A row the app never touches keeps every cell it had."""
    plan = build_plan(visits, schedule_rows)
    worksheet, _ = applied_sheet(schedule_bytes, plan)

    delacroix = next(r for r in schedule_rows if str(r.patient).startswith("Delacroix"))
    found = find_row(worksheet, "Delacroix, Owen")
    assert found["Payment"] == delacroix.payment
    assert found["DX"] == delacroix.dx
    assert found["Billed"] == delacroix.billed


def test_relative_order_of_existing_rows_is_preserved(schedule_bytes, visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    worksheet, _ = applied_sheet(schedule_bytes, plan)

    original_order = [str(r.patient) for r in schedule_rows]
    final = patients_in_order(worksheet)
    # Every original name still appears, in the same relative sequence.
    positions = []
    search_from = 0
    for name in original_order:
        index = final.index(name, search_from)
        positions.append(index)
        search_from = index + 1
    assert positions == sorted(positions)


# --- Styling ---------------------------------------------------------------

def test_inserted_rows_inherit_the_anchor_row_styling(schedule_bytes, visits, schedule_rows):
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    worksheet, _ = applied_sheet(schedule_bytes, build_plan([visit], schedule_rows))

    for column in (PATIENT_COL, DATA_COL, PAYMENT_COL):
        anchor = worksheet.cell(row=9, column=column)
        inserted = worksheet.cell(row=10, column=column)
        assert inserted.number_format == anchor.number_format
        assert inserted.font.name == anchor.font.name
        assert inserted.font.sz == anchor.font.sz


def test_inserted_payment_cell_keeps_the_numeric_format(schedule_bytes, visits, schedule_rows):
    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    worksheet, _ = applied_sheet(schedule_bytes, build_plan([visit], schedule_rows))
    assert worksheet.cell(row=10, column=PAYMENT_COL).number_format == "0.00"


# --- Structural safety -----------------------------------------------------

def _sheet_with(schedule_bytes, mutate):
    workbook = load_schedule_workbook(schedule_bytes)
    mutate(workbook[SHEET_NAME])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_merged_cells_below_an_anchor_block_insertion(schedule_bytes):
    """Splitting a merged range would corrupt the sheet, so we append instead."""
    data = _sheet_with(schedule_bytes, lambda ws: ws.merge_cells("A9:B10"))
    workbook = load_schedule_workbook(data)
    worksheet = get_schedule_sheet(workbook)
    columns = resolve_columns(worksheet)

    assert insertion_blocked_reason(worksheet, columns, 9) is not None
    assert "merged" in insertion_blocked_reason(worksheet, columns, 9)


def test_totals_row_below_an_anchor_blocks_insertion(schedule_bytes):
    def add_totals(ws):
        ws.insert_rows(10)
        ws.cell(row=10, column=2, value="TOTALS")
        ws.cell(row=10, column=5, value=999.0)

    data = _sheet_with(schedule_bytes, add_totals)
    workbook = load_schedule_workbook(data)
    worksheet = get_schedule_sheet(workbook)
    columns = resolve_columns(worksheet)

    reason = insertion_blocked_reason(worksheet, columns, 9)
    assert reason is not None and "totals" in reason


def test_blocked_anchor_falls_back_to_appending(schedule_bytes, visits):
    """The row still gets written -- at the bottom -- and is flagged."""
    data = _sheet_with(schedule_bytes, lambda ws: ws.merge_cells("A9:B10"))
    workbook = load_schedule_workbook(data)
    worksheet = get_schedule_sheet(workbook)
    rows = read_schedule_rows(worksheet, resolve_columns(worksheet))

    visit = find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 4, 9))
    plan = build_plan([visit], rows)

    updated, stats = build_updated_workbook(data, plan)
    assert stats["inserted_rows"] == 0
    assert stats["appended_rows"] == 1
    assert plan[0].placement_fallback
    assert "appended" in plan[0].placement_display

    out = load_schedule_workbook(updated.getvalue())[SHEET_NAME]
    assert out.cell(row=out.max_row, column=PATIENT_COL).value == "Castellano, Miguel"


def test_a_clean_anchor_is_not_blocked(schedule_bytes):
    workbook = load_schedule_workbook(schedule_bytes)
    worksheet = get_schedule_sheet(workbook)
    columns = resolve_columns(worksheet)
    assert insertion_blocked_reason(worksheet, columns, 9) is None


# --- Preview ---------------------------------------------------------------

def test_placement_display_strings(visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    grouped = next(c for c in plan
                   if c.visit.patient == "CASTELLANO, MIGUEL"
                   and c.effective_action == ACTION_NEW)
    appended = next(c for c in plan if c.visit.patient == "WHITCOMBE, ROSALIND")

    assert grouped.placement_display == 'under existing "Castellano, Miguel"'
    assert appended.placement_display == "appended (patient not in schedule)"


def test_placement_display_is_blank_for_non_new_rows(visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    skipped = next(c for c in plan if c.action == "Skip (already paid)")
    assert skipped.placement_display == ""


# --- Idempotency still holds ------------------------------------------------

def test_second_run_over_a_grouped_sheet_is_a_no_op(schedule_bytes, visits, schedule_rows):
    plan = build_plan(visits, schedule_rows)
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    updated_bytes = updated.getvalue()

    worksheet = get_schedule_sheet(load_schedule_workbook(updated_bytes))
    rows = read_schedule_rows(worksheet, resolve_columns(worksheet))
    second = build_plan(visits, rows)

    _, stats = build_updated_workbook(updated_bytes, second)
    assert stats["new_rows"] == 0
    assert stats["filled_cells"] == 0

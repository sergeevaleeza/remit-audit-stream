"""Regenerate `tests/fixtures/List_of_Patients_Schedule.xlsx`.

Every patient below is fictional (see `tests/fixtures/synthetic_remit_data.py`,
the shared source of truth with the synthetic remittance PDF). The sample
deliberately mirrors the quirks of a real clinic's workbook: a title row above
the headers, `Data`/`Billed` values that are a mix of real datetimes and
hand-typed strings, junk placeholders, a formula in `Co-pay`, and a `0.00`
payment -- plus a second sheet that must survive the round trip untouched.

Run with:  python -m tests.make_fixture
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .fixtures.synthetic_employees_data import (
    ANA_HEADERS,
    ANA_ROWS,
    MARCIA_HEADERS,
    MARCIA_ROWS,
    OXANA_HEADERS,
    OXANA_ROWS,
    OXANA_TITLE,
)
from .fixtures.synthetic_remit_data import MUTUAL_ROWS, SCHEDULE_ONLY_PATIENT

FIXTURE = Path(__file__).parent / "fixtures" / "List_of_Patients_Schedule.xlsx"
MUTUAL_FIXTURE = Path(__file__).parent / "fixtures" / "List_of_Patients_Mutual.xlsx"
EMPLOYEES_FIXTURE = Path(__file__).parent / "fixtures" / "AMSMC_employees_sample.xlsx"

HEADERS = [
    "Patient", "Ins", "Data", "Billed", "Payment", "Co-pay",
    "Co-pays Paid", "Office", "Comment", "DX", "CPT Code",
]

# A fictional diagnosis set, invented for these fictional patients -- not
# copied from any real chart.
DX = "F41.1, F32.9"

DR_A = "Dr. A"

# Patient, Ins, Data, Billed, Payment, Co-pay, Co-pays Paid, Office, Comment, DX, CPT
ROWS = [
    # --- Already paid: must stay untouched (idempotency) --------------------
    # This Comment names a provider *tab*, so the employees-file cross-check
    # has something to fire on. A `Dr. …` comment names a physician, not a tab,
    # and is deliberately not treated as a conflict.
    ["Marlowe, Diane", "Medicare", datetime(2026, 3, 9), "2/32/26",
     166.37, "=23.52+18.92", None, None, "Oxana", DX, "99213/90833"],
    ["Marlowe, Diane", "Medicare", "04/21/2026", datetime(2026, 6, 18),
     224.59, 57.29, None, None, None, DX, "99214/90836"],
    ["Thackeray, Renata", "Medicare", datetime(2026, 3, 20), "2/32/26",
     166.37, "=23.52+18.92", None, None, DR_A, DX, "99213/90833"],
    ["Thackeray, Renata", "Medicare", "4/7/2026", datetime(2026, 6, 18),
     186.31, 47.53, None, None, None, DX, "99213/90836"],
    ["Thackeray, Renata", "Medicare", datetime(2026, 4, 24), "No Billing",
     166.37, 42.44, None, None, DR_A, DX, "99213/90833"],

    # --- Zero payment is a recorded value, not a blank ----------------------
    ["Whitfield, Harold", "Medicare", datetime(2026, 3, 4), "2/32/26",
     0.00, None, None, None, None, DX, "99214"],

    # --- Fill path: empty Payment, prefilled Billed placeholder + DX/CPT ----
    ["Castellano, Miguel", "Medicare", datetime(2026, 3, 26), "2/32/26",
     None, None, None, None, None, DX, "99213/90836"],

    # --- Medicare-truncated surname in the remit ----------------------------
    ["Featherstonehaugh, Wilhelmina", "Medicare", "3/2/2026", "2/32/26",
     None, None, None, None, None, DX, "99213/90833"],

    # --- Trailing middle initial in the remit -------------------------------
    ["Sorensen, Marcus", "Medicare", datetime(2026, 5, 14), "2/32/26",
     None, None, None, None, None, DX, "99213/90836"],

    # --- Generational suffix in the sheet, absent from the remit ------------
    # DX is left blank here so the Mutual-file fill path has something to do.
    ["Bystritskaya Jr, Anna", "Medicare", datetime(2026, 4, 14), "2/32/26",
     None, None, None, None, None, None, "99213/90833"],

    # --- Date present in the sheet but not in the remit ---------------------
    [SCHEDULE_ONLY_PATIENT["schedule_name"], "Medicare", datetime(2026, 1, 5), "2/32/26",
     None, None, None, None, None, DX, "99213/90833"],
]


def build() -> Workbook:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "2026 Medicare"

    title = sheet.cell(row=1, column=1, value=2026)
    title.font = Font(bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=2, column=index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for offset, values in enumerate(ROWS):
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=3 + offset, column=index, value=value)
            if index in (3, 4) and isinstance(value, datetime):
                cell.number_format = "mm/dd/yyyy"
            if index in (5, 6):
                cell.number_format = "0.00"

    # A second sheet that must come back byte-for-byte intact.
    other = workbook.create_sheet("2026 Medical")
    other["A1"] = "Untouched sheet"
    other["A2"] = "Patient"
    other["B2"] = "Note"
    other["A3"] = "Example, Patient"
    other["B3"] = "must survive the round trip"
    other["C3"] = "=1+1"
    other["A1"].font = Font(bold=True, italic=True)

    return workbook


def build_mutual() -> Workbook:
    """The optional DX reference workbook: sheet `Active`, no header row.

    Column A = Patient, B = DX, E = attending doctor. C, D and E exist only
    for structural realism -- the app reads A and B and nothing else.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Active"

    for index, (patient, dx, attending) in enumerate(MUTUAL_ROWS, start=1):
        sheet.cell(row=index, column=1, value=patient)
        sheet.cell(row=index, column=2, value=dx)
        sheet.cell(row=index, column=3, value="unused")
        sheet.cell(row=index, column=4, value="unused")
        sheet.cell(row=index, column=5, value=attending)

    # A second sheet the app must ignore entirely.
    inactive = workbook.create_sheet("Inactive")
    inactive["A1"] = "Former, Patient"
    inactive["B1"] = "F99.9"

    return workbook


def _write_employee_sheet(sheet, headers, rows, header_row: int) -> None:
    """Lay out one provider tab: headers on `header_row`, data below."""
    header_fill = PatternFill("solid", fgColor="E2EFDA")
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    lookup = {h: i for i, h in enumerate(headers, start=1)}
    for offset, (patient, session_date, prefilled) in enumerate(rows):
        row = header_row + 1 + offset
        sheet.cell(row=row, column=lookup[headers[0]], value=patient)
        sheet.cell(row=row, column=lookup["Date of Session"], value=session_date)
        for label, value in prefilled.items():
            cell = sheet.cell(row=row, column=lookup[label], value=value)
            cell.number_format = "0.00"


def build_employees() -> Workbook:
    """The synthetic `AMSMC_employees.xlsx`: three tabs, two header layouts."""
    workbook = Workbook()

    ana = workbook.active
    ana.title = "Ana"
    _write_employee_sheet(ana, ANA_HEADERS, ANA_ROWS, header_row=1)

    marcia = workbook.create_sheet("Marcia")
    _write_employee_sheet(marcia, MARCIA_HEADERS, MARCIA_ROWS, header_row=1)

    # Oxana's sheet carries a title on row 1, so her headers sit on row 2.
    oxana = workbook.create_sheet("Oxana")
    oxana.cell(row=1, column=1, value=OXANA_TITLE).font = Font(bold=True, size=14)
    _write_employee_sheet(oxana, OXANA_HEADERS, OXANA_ROWS, header_row=2)

    # A sheet the app must ignore entirely.
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Not a provider tab"

    return workbook


def main() -> None:
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    build().save(FIXTURE)
    print(f"wrote {FIXTURE}")
    build_mutual().save(MUTUAL_FIXTURE)
    print(f"wrote {MUTUAL_FIXTURE}")
    build_employees().save(EMPLOYEES_FIXTURE)
    print(f"wrote {EMPLOYEES_FIXTURE}")


if __name__ == "__main__":
    main()

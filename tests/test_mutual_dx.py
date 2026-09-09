"""The optional Mutual DX reference upload (Change 3)."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from remit.config import (
    ACTION_FILL,
    ACTION_NEW,
    ACTION_REVIEW,
    ACTION_SKIP,
    COL_DX,
    DX_AMBIGUOUS,
    DX_CONFLICT,
    DX_FOUND,
    DX_NO_FILE,
    DX_NOT_FOUND,
    SHEET_NAME,
)
from remit.excel_updater import (
    ScheduleError,
    build_updated_workbook,
    load_schedule_workbook,
)
from remit.matching import build_plan, plan_change
from remit.mutual import build_dx_lookup, load_dx_lookup

from .conftest import find_visit

MUTUAL_XLSX = Path(__file__).parent / "fixtures" / "List_of_Patients_Mutual.xlsx"


@pytest.fixture()
def mutual_bytes() -> bytes:
    return MUTUAL_XLSX.read_bytes()


@pytest.fixture()
def dx_lookup(mutual_bytes):
    return load_dx_lookup(mutual_bytes)


def change_for(plan, patient: str, service_date: date):
    for change in plan:
        if change.visit.patient == patient and change.visit.service_date == service_date:
            return change
    raise AssertionError(f"no planned change for {patient} on {service_date}")


# --- Reading the file ------------------------------------------------------

def test_reads_the_active_sheet_positionally(dx_lookup):
    """No header row: row 1 is data, column A is the patient, B the DX."""
    assert len(dx_lookup) == 10
    assert dx_lookup.find("MARLOWE, DIANE").dx == "F41.1, F32.9"


def test_ignores_other_sheets(dx_lookup):
    """The `Inactive` sheet must not contribute entries."""
    assert dx_lookup.find("FORMER, PATIENT").status == DX_NOT_FOUND


def test_column_e_attending_doctor_is_not_read(mutual_bytes):
    """Column E exists in the fixture but must never influence anything."""
    sheet = openpyxl.load_workbook(io.BytesIO(mutual_bytes))["Active"]
    assert sheet.cell(row=1, column=5).value == "Dr. A"  # present in the file

    lookup = load_dx_lookup(mutual_bytes)
    for entry in lookup.entries:
        assert "Dr." not in entry.dx


def test_missing_active_sheet_is_reported_clearly():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Something Else"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ScheduleError, match="Active"):
        load_dx_lookup(buffer.getvalue())


def test_rows_without_a_dx_are_skipped():
    lookup = build_dx_lookup([("Someone, Real", "F41.1", 1), ("Blank, Person", "", 2)])
    assert len(lookup) == 1
    assert lookup.find("BLANK, PERSON").status == DX_NOT_FOUND


# --- Name matching into the lookup -----------------------------------------

def test_uppercase_remit_name_matches_title_case_mutual_name(dx_lookup):
    assert dx_lookup.find("MARLOWE, DIANE").status == DX_FOUND


def test_medicare_truncated_name_matches(dx_lookup):
    """`FEATHERSTONEH, WILHELMI` must find `Featherstonehaugh, Wilhelmina`."""
    match = dx_lookup.find("FEATHERSTONEH, WILHELMI")
    assert match.status == DX_FOUND
    assert match.dx == "F31.81"


def test_trailing_middle_initial_matches(dx_lookup):
    match = dx_lookup.find("SORENSEN, MARCUS R")
    assert match.status == DX_FOUND
    assert match.dx == "F40.10"


def test_generational_suffix_matches_across_files(dx_lookup):
    """Change-1 normalisation applies to the Mutual lookup too."""
    match = dx_lookup.find("BYSTRITSKAYA, ANNA")
    assert match.status == DX_FOUND
    assert match.dx == "F42.2"


def test_patient_absent_from_mutual_is_not_found(dx_lookup):
    assert dx_lookup.find("OKAFOR, CHIDINMA").status == DX_NOT_FOUND
    assert dx_lookup.find("OKAFOR, CHIDINMA").dx is None


def test_conflicting_dx_uses_the_first_and_reports_it(dx_lookup):
    match = dx_lookup.find("WHITCOMBE, ROSALIND")
    assert match.status == DX_CONFLICT
    assert match.dx == "F34.1"                       # first wins
    assert match.conflicting == ("F34.1", "F33.2")


def test_conflicts_are_listed_on_the_lookup(dx_lookup):
    assert dx_lookup.conflicts == {"whitcombe|rosalind": ("F34.1", "F33.2")}


def test_ambiguous_name_is_flagged_not_guessed():
    """A near-miss name in the review band must not silently pick a DX."""
    lookup = build_dx_lookup([("Testname, Sammpel", "F99.1", 1)])
    match = lookup.find("TESTNAME, SAMPLE")
    assert match.status == DX_AMBIGUOUS
    assert match.dx is None


def test_unrelated_name_is_not_found():
    lookup = build_dx_lookup([("Zimmerman, Robert", "F99.1", 1)])
    assert lookup.find("MARLOWE, DIANE").status == DX_NOT_FOUND


# --- Planning with DX ------------------------------------------------------

def test_new_row_gets_dx_from_mutual(visits, schedule_rows, dx_lookup):
    change = plan_change(
        find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6)),
        schedule_rows, dx_lookup,
    )
    assert change.action == ACTION_NEW
    assert change.dx_value == "F34.1"


def test_new_row_without_mutual_leaves_dx_blank(visits, schedule_rows):
    change = plan_change(
        find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6)), schedule_rows
    )
    assert change.dx_value is None
    assert change.dx_status == DX_NO_FILE


def test_existing_row_with_blank_dx_is_filled(visits, schedule_rows, dx_lookup):
    """`Bystritskaya Jr, Anna` has a blank DX in the fixture."""
    change = plan_change(
        find_visit(visits, "BYSTRITSKAYA, ANNA", date(2026, 4, 14)),
        schedule_rows, dx_lookup,
    )
    assert change.action == ACTION_FILL
    assert change.fills[COL_DX] == "F42.2"


def test_existing_non_blank_dx_is_never_overwritten(visits, schedule_rows, dx_lookup):
    """`Castellano, Miguel` already has a DX, so it must be left alone."""
    change = plan_change(
        find_visit(visits, "CASTELLANO, MIGUEL", date(2026, 3, 26)),
        schedule_rows, dx_lookup,
    )
    assert change.action == ACTION_FILL
    assert COL_DX not in change.fills


def test_skipped_rows_are_not_given_a_dx(visits, schedule_rows, dx_lookup):
    """An already-paid row stays completely untouched, DX included."""
    change = plan_change(
        find_visit(visits, "MARLOWE, DIANE", date(2026, 3, 9)),
        schedule_rows, dx_lookup,
    )
    assert change.action == ACTION_SKIP
    assert change.fills == {}


def test_ambiguous_dx_match_needs_review(visits, schedule_rows):
    """A ~88% name match is too weak to pick a diagnosis from."""
    lookup = build_dx_lookup([("Whitcambe, Rosalynd", "F99.1", 1)])
    change = plan_change(
        find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6)),
        schedule_rows, lookup,
    )
    assert change.action == ACTION_REVIEW
    assert change.dx_value is None
    assert any("Ambiguous DX" in reason for reason in change.review_reasons)


# --- Writing DX ------------------------------------------------------------

def test_dx_is_written_to_the_workbook(schedule_bytes, visits, schedule_rows, dx_lookup):
    plan = build_plan(visits, schedule_rows, dx_lookup)
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    bystritskaya = next(
        r for r in schedule_rows if str(r.patient).startswith("Bystritskaya")
    )
    assert worksheet.cell(row=bystritskaya.row_num, column=10).value == "F42.2"


def test_new_row_dx_lands_in_the_dx_column(schedule_bytes, visits, schedule_rows, dx_lookup):
    visit = find_visit(visits, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    plan = build_plan([visit], schedule_rows, dx_lookup)
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    row = len(schedule_rows) + 3
    assert worksheet.cell(row=row, column=1).value == "Whitcombe, Rosalind"
    assert worksheet.cell(row=row, column=10).value == "F34.1"


def test_not_found_patient_leaves_dx_blank(schedule_bytes, visits, schedule_rows, dx_lookup):
    visit = find_visit(visits, "OKAFOR, CHIDINMA", date(2026, 4, 28))
    plan = build_plan([visit], schedule_rows, dx_lookup)
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    row = len(schedule_rows) + 3
    assert worksheet.cell(row=row, column=1).value == "Okafor, Chidinma"
    assert worksheet.cell(row=row, column=10).value is None


def test_without_mutual_behaviour_is_unchanged(schedule_bytes, visits, schedule_rows):
    """The upload is optional: no file means DX blank, exactly as before."""
    plan = build_plan(visits, schedule_rows)
    updated, _ = build_updated_workbook(schedule_bytes, plan)
    worksheet = load_schedule_workbook(updated.getvalue())[SHEET_NAME]

    for offset in range(len(schedule_rows) + 3, worksheet.max_row + 1):
        assert worksheet.cell(row=offset, column=10).value is None


def test_dx_display_strings(visits, schedule_rows, dx_lookup):
    plan = build_plan(visits, schedule_rows, dx_lookup)
    conflict = change_for(plan, "WHITCOMBE, ROSALIND", date(2026, 5, 6))
    missing = change_for(plan, "OKAFOR, CHIDINMA", date(2026, 4, 28))
    found = change_for(plan, "BYSTRITSKAYA, ANNA", date(2026, 4, 14))

    assert conflict.dx_display == "F34.1 (conflict: F34.1, F33.2)"
    assert missing.dx_display == f"({DX_NOT_FOUND})"
    assert found.dx_display == "F42.2"
